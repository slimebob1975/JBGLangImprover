import json
import logging
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from lxml import etree

from app.src import JBGLangImprovSuggestorAI as suggestor_module
from app.src.JBGContentClassifier import generated_kind_for_field
from app.src.JBGDocumentStructureExtractor import DocumentStructureExtractor
from app.src.JBGGlobalAnalyzerAI import GlobalReviewResult, JBGGlobalAnalyzerAI, build_outline
from app.src.JBGLangImprovSuggestorAI import JBGLangImprovSuggestorAI
from app.src.JBGLanguageImprover import JBGLanguageImprover
from app.src.JBGReadabilityMetrics import compute_document_readability


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


def _quiet_logger(name):
    logger = logging.getLogger(name)
    logger.handlers.clear()
    logger.addHandler(logging.NullHandler())
    logger.propagate = False
    return logger


def _xml(fragment):
    return parse_xml(fragment.replace("<w:", f"<w:", 1).replace(">", f" {nsdecls('w')}>", 1))


def _texts(structure):
    return {e["element_id"]: e["text"] for e in structure["elements"]}


class _Base(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"extractor-{id(self)}")

    def tearDown(self):
        self.temp_dir.cleanup()

    def extract(self, document):
        path = self.root / f"doc_{id(document)}.docx"
        document.save(path)
        return path, DocumentStructureExtractor(str(path), self.logger).extract()


# ============================================================================
# 1. Väntande infogningar från andra skribenter
# ============================================================================

def build_pending_changes_document(path):
    """Testkörningens fall: Annas infogning, borttagning och kommentarer."""
    document = Document()
    p1 = document.add_paragraph("Myndigheten har genomfört en granskning av ärendet.")
    p2 = document.add_paragraph("Rapporten beskriver ")
    p2._p.append(_xml('<w:ins w:id="7" w:author="Anna" w:date="2026-01-01T00:00:00Z">'
                      '<w:r><w:t xml:space="preserve">utförligt </w:t></w:r></w:ins>'))
    p2._p.append(_xml('<w:del w:id="8" w:author="Anna" w:date="2026-01-01T00:00:00Z">'
                      '<w:r><w:delText xml:space="preserve">kortfattat </w:delText></w:r></w:del>'))
    p2.add_run("hur genomförandet av kontrollerna gick till.")
    document.add_comment(p1.runs, text="Annas kommentar", author="Anna", initials="A")
    document.save(path)


class PendingInsertionTests(_Base):
    def test_text_includes_pending_insertions_and_skips_pending_deletions(self):
        path = self.root / "pending.docx"
        build_pending_changes_document(path)
        texts = _texts(DocumentStructureExtractor(str(path), self.logger).extract())
        self.assertEqual(texts["paragraph_2"], "Rapporten beskriver utförligt hur genomförandet av kontrollerna gick till.")

    def test_ordinary_paragraphs_keep_exactly_the_same_text(self):
        # Utan spårade ändringar ska texten vara densamma som python-docx gav förut
        document = Document()
        document.add_paragraph("Vanlig text med\ttabb.")
        with_break = document.add_paragraph("Rad ett")
        with_break.add_run().add_break()
        with_break.add_run("rad två.")
        document.add_paragraph("Rubrik", style="Heading 1")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Cell med text."
        path, structure = self.extract(document)
        reference = Document(str(path))
        texts = _texts(structure)
        for index, para in enumerate(reference.paragraphs, start=1):
            self.assertEqual(texts[f"paragraph_{index}"], para.text)
        self.assertEqual(texts["table_1_cell_1_1_p1"], reference.tables[0].cell(0, 0).paragraphs[0].text)

    def test_proposal_after_a_pending_insertion_is_applied_and_old_markup_kept(self):
        source = self.root / "pending.docx"
        build_pending_changes_document(source)
        content = json.dumps([{
            "type": "paragraph", "element_id": "paragraph_2",
            "old": "genomförandet av kontrollerna gick till", "new": "kontrollerna genomfördes",
            "motivation": "Kortare.",
        }], ensure_ascii=False)
        client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **k: SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))], usage=None))))
        for mode in ("tracked", "simple"):
            with self.subTest(mode=mode):
                output = self.root / f"out_{mode}.docx"
                with mock.patch.object(suggestor_module.openai, "OpenAI", return_value=client):
                    improver = JBGLanguageImprover(
                        input_path=str(source), api_key="k", model="m", prompt_policy="p", temperature=1,
                        include_motivations=True, logger=self.logger, docx_mode=mode,
                        include_about_section=False, compute_readability=False,
                    )
                    improver.run(output_path=str(output))
                counts = improver.run_summary.local_suggestions
                self.assertEqual((counts.applied, counts.failed), (1, 0))
                with ZipFile(output) as z:
                    body = etree.fromstring(z.read("word/document.xml"))
                    comments = etree.fromstring(z.read("word/comments.xml"))
                annas = [e for e in body.iter(f"{W}ins", f"{W}del") if e.get(f"{W}author") == "Anna"]
                self.assertEqual(len(annas), 2)
                self.assertIn("Anna", [c.get(f"{W}author") for c in comments.findall(f"{W}comment")])


# ============================================================================
# 2. Helt dold text
# ============================================================================

class HiddenTextTests(_Base):
    def build(self):
        document = Document()
        hidden_paragraph_style = document.styles.add_style("Mallinstruktion", WD_STYLE_TYPE.PARAGRAPH)
        hidden_paragraph_style.font.hidden = True
        hidden_character_style = document.styles.add_style("Dold text", WD_STYLE_TYPE.CHARACTER)
        hidden_character_style.font.hidden = True

        document.add_paragraph("Synligt första stycke.")                                   # paragraph_1
        direct = document.add_paragraph()                                                  # paragraph_2
        direct.add_run("Ta ej bort denna avsnittsbrytning!!").font.hidden = True
        document.add_paragraph("Mallens instruktion.", style="Mallinstruktion")            # paragraph_3
        by_character_style = document.add_paragraph()                                      # paragraph_4
        by_character_style.add_run("Dold via teckenformat.", style="Dold text")
        overridden = document.add_paragraph(style="Mallinstruktion")                       # paragraph_5
        overridden.add_run("Synlig trots formatmallen.").font.hidden = False
        partly = document.add_paragraph("Delvis synligt ")                                 # paragraph_6
        partly.add_run("och delvis dolt.").font.hidden = True
        document.add_paragraph("Sista synliga stycket.")                                   # paragraph_7
        cell = document.add_table(rows=1, cols=2)
        cell.cell(0, 0).paragraphs[0].add_run("Dold cell.").font.hidden = True
        cell.cell(0, 1).text = "Synlig cell."
        return document

    def test_fully_hidden_paragraphs_are_left_out_without_shifting_ids(self):
        _, structure = self.extract(self.build())
        texts = _texts(structure)
        for hidden_id in ("paragraph_2", "paragraph_3", "paragraph_4", "table_1_cell_1_1_p1"):
            self.assertNotIn(hidden_id, texts)
        self.assertEqual(texts["paragraph_1"], "Synligt första stycke.")
        self.assertEqual(texts["paragraph_5"], "Synlig trots formatmallen.")
        self.assertEqual(texts["paragraph_7"], "Sista synliga stycket.")
        self.assertEqual(texts["table_1_cell_1_2_p1"], "Synlig cell.")
        self.assertEqual(sorted(structure["excluded"]["hidden"]),
                         ["paragraph_2", "paragraph_3", "paragraph_4", "table_1_cell_1_1_p1"])

    def test_partly_hidden_paragraph_is_kept_unchanged(self):
        # Att ta bort enstaka dolda ord skulle flytta förankringen av förslagen
        _, structure = self.extract(self.build())
        self.assertEqual(_texts(structure)["paragraph_6"], "Delvis synligt och delvis dolt.")

    def test_hidden_text_is_not_counted_in_lix(self):
        _, structure = self.extract(self.build())
        words = compute_document_readability(structure, []).before.words
        self.assertNotIn("avsnittsbrytning", json.dumps(structure, ensure_ascii=False))
        self.assertEqual(words, 3 + 3 + 5 + 3 + 2)   # stycke 1, 5, 6, 7 och den synliga cellen


# ============================================================================
# 3. Automatiskt genererat innehåll
# ============================================================================

def _toc_field_paragraphs(document):
    """En innehållsförteckning som fält över tre stycken, med PAGEREF inuti."""
    first = document.add_paragraph()
    first._p.append(_xml('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'))
    first._p.append(_xml('<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'))
    first._p.append(_xml('<w:r><w:fldChar w:fldCharType="separate"/></w:r>'))
    first.add_run("Inledning\t3")
    second = document.add_paragraph()
    second._p.append(_xml('<w:r><w:fldChar w:fldCharType="begin"/></w:r>'))
    second._p.append(_xml('<w:r><w:instrText xml:space="preserve"> PAGEREF _Toc1 \\h </w:instrText></w:r>'))
    second._p.append(_xml('<w:r><w:fldChar w:fldCharType="separate"/></w:r>'))
    second.add_run("Resultat\t5")
    second._p.append(_xml('<w:r><w:fldChar w:fldCharType="end"/></w:r>'))
    third = document.add_paragraph()
    third._p.append(_xml('<w:r><w:fldChar w:fldCharType="end"/></w:r>'))


class GeneratedContentTests(_Base):
    def test_toc_field_over_several_paragraphs_becomes_one_placeholder(self):
        document = Document()
        document.add_paragraph("Innehåll", style="Heading 1")      # paragraph_1
        _toc_field_paragraphs(document)                             # paragraph_2-4
        document.add_paragraph("Inledning", style="Heading 1")     # paragraph_5
        document.add_paragraph("Text efter förteckningen.")        # paragraph_6
        _, structure = self.extract(document)

        ordered = sorted(structure["elements"], key=lambda e: e["doc_order"] or 0)
        self.assertEqual([e["element_id"] for e in ordered],
                         ["paragraph_1", "generated_1", "paragraph_4", "paragraph_5", "paragraph_6"])
        placeholder = ordered[1]
        self.assertEqual((placeholder["type"], placeholder["generated_kind"]), ("generated", "toc"))
        self.assertEqual(placeholder["text"], "[Innehållsförteckning, skapas automatiskt]")
        self.assertEqual(structure["excluded"]["generated"], ["paragraph_2", "paragraph_3"])
        # Ett tomt stycke med bara fältets slut räknas inte som förteckning
        self.assertEqual(_texts(structure)["paragraph_4"], "")

    def test_toc_in_a_content_control_becomes_a_placeholder_in_reading_order(self):
        # Testkörningen 905c7230: "Innehåll" följdes av en förteckning i en innehållskontroll
        document = Document()
        document.add_paragraph("Innehåll", style="Heading 1")
        sdt = _xml('<w:sdt><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
                   '<w:docPartUnique/></w:docPartObj></w:sdtPr><w:sdtContent>'
                   '<w:p><w:r><w:t>Inledning 3</w:t></w:r></w:p></w:sdtContent></w:sdt>')
        document.element.body.sectPr.addprevious(sdt)
        document.add_paragraph("Inledning", style="Heading 1")
        _, structure = self.extract(document)

        outline = build_outline(structure)
        self.assertEqual(outline, [
            {"id": "paragraph_1", "h": 1, "t": "Innehåll"},
            {"id": "generated_1", "t": "[Innehållsförteckning, skapas automatiskt]"},
            {"id": "paragraph_2", "h": 1, "t": "Inledning"},
        ])

    def test_toc_styled_paragraphs_without_field_are_recognized(self):
        document = Document()
        document.styles.add_style("toc 1", WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph("Innehåll", style="Heading 1")
        document.add_paragraph("Inledning\t3", style="toc 1")
        document.add_paragraph("Resultat\t5", style="toc 1")
        document.add_paragraph("Brödtext.")
        _, structure = self.extract(document)
        generated = [e for e in structure["elements"] if e["type"] == "generated"]
        self.assertEqual(len(generated), 1)
        self.assertEqual(structure["excluded"]["generated"], ["paragraph_2", "paragraph_3"])

    def test_field_kinds(self):
        self.assertEqual(generated_kind_for_field(' TOC \\o "1-3" \\h '), "toc")
        self.assertEqual(generated_kind_for_field(' TOC \\h \\z \\c "Figur" '), "figures")
        self.assertEqual(generated_kind_for_field(" INDEX \\e "), "index")
        self.assertEqual(generated_kind_for_field(" BIBLIOGRAPHY "), "bibliography")
        self.assertIsNone(generated_kind_for_field(" PAGEREF _Toc1 \\h "))

    def test_placeholders_are_never_reviewed_locally(self):
        document = Document()
        document.add_paragraph("Innehåll", style="Heading 1")
        _toc_field_paragraphs(document)
        path, structure = self.extract(document)
        structure_path = self.root / "s.json"
        structure_path.write_text(json.dumps(structure, ensure_ascii=False), encoding="utf-8")
        suggestor = JBGLangImprovSuggestorAI(api_key="k", model="m", prompt_policy="p",
                                             temperature=1, logger=self.logger)
        suggestor.load_structure(str(structure_path))
        types = {e["type"] for e in suggestor.json_structured_document["elements"]}
        self.assertNotIn("generated", types)
        self.assertNotIn("excluded", suggestor.json_structured_document)

    def test_global_findings_can_never_involve_a_placeholder(self):
        structure = {"elements": [
            {"type": "paragraph", "element_id": "paragraph_1", "text": "Innehåll", "heading_level": 1, "doc_order": 1},
            {"type": "generated", "element_id": "generated_1", "text": "[Innehållsförteckning, skapas automatiskt]",
             "doc_order": 2, "generated_kind": "toc"},
            {"type": "paragraph", "element_id": "paragraph_3", "text": "Text.", "doc_order": 3},
        ]}
        analyzer = JBGGlobalAnalyzerAI("k", "m", 1, self.logger, policy="p")
        result = GlobalReviewResult()
        analyzer.validate(structure, [(1, {
            "category": "error", "element_ids": ["generated_1"],
            "quote": "Innehållsförteckning", "description": "Tom förteckning.",
        }), (1, {
            "category": "repetition", "element_ids": ["generated_1", "paragraph_3"],
            "quote": "Text.", "description": "Upprepning.",
        })], result)
        self.assertEqual([r.reason for r in result.rejected],
                         ["generated_content", "repetition_needs_two_locations"])


if __name__ == "__main__":
    unittest.main()
