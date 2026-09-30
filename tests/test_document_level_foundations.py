import json
import logging
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.src import JBGLangImprovSuggestorAI as suggestor_module
from app.src.JBGDocumentStructureExtractor import DocumentStructureExtractor
from app.src.JBGLanguageImprover import JBGLanguageImprover
from app.src.JBGUsageTracker import UsageTracker


def _quiet_logger(name):
    logger = logging.getLogger(name)
    logger.handlers.clear()
    logger.addHandler(logging.NullHandler())
    logger.propagate = False
    return logger


def _set_outline_level(p_pr_owner, level):
    p_pr = p_pr_owner.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


class DocumentMetadataTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"metadata-test-{id(self)}")
        self.source = self.root / "source.docx"

        document = Document()
        document.add_paragraph("Rapportens titel", style="Title")          # paragraph_1
        document.add_paragraph("Inledning", style="Heading 1")             # paragraph_2
        document.add_paragraph("Vanlig brödtext.")                         # paragraph_3

        # Egen stil med svenskt style-id och outlineLvl 1 -> rubriknivå 2
        swedish = document.styles.add_style("Rubrik2", WD_STYLE_TYPE.PARAGRAPH)
        _set_outline_level(swedish.element, 1)
        document.add_paragraph("Bakgrund", style="Rubrik2")                # paragraph_4

        # Egen stil utan outlineLvl som ärver från Heading 1
        inherited = document.styles.add_style("Kapitelrubrik", WD_STYLE_TYPE.PARAGRAPH)
        inherited.base_style = document.styles["Heading 1"]
        document.add_paragraph("Resultat", style="Kapitelrubrik")          # paragraph_5

        table = document.add_table(rows=1, cols=2)                         # table_1
        table.cell(0, 0).text = "Cell A"
        table.cell(0, 1).text = "Cell B"

        direct = document.add_paragraph("Direktformaterad rubrik")         # paragraph_6
        _set_outline_level(direct._p, 2)

        document.add_paragraph("Avslutande text.")                         # paragraph_7
        document.save(self.source)

        self.structure = DocumentStructureExtractor(str(self.source), self.logger).extract()
        self.by_id = {e["element_id"]: e for e in self.structure["elements"]}

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_existing_element_ids_are_unchanged(self):
        ids = [e["element_id"] for e in self.structure["elements"]]
        self.assertEqual(ids[:7], [f"paragraph_{i}" for i in range(1, 8)])
        self.assertIn("table_1_cell_1_1_p1", ids)
        self.assertIn("table_1_cell_1_2_p1", ids)

    def test_heading_levels(self):
        levels = {i: self.by_id[f"paragraph_{i}"]["heading_level"] for i in range(1, 8)}
        self.assertEqual(levels, {1: 0, 2: 1, 3: None, 4: 2, 5: 1, 6: 3, 7: None})

    def test_style_ids_and_names(self):
        self.assertEqual(self.by_id["paragraph_2"]["style_id"], "Heading1")
        self.assertEqual(self.by_id["paragraph_4"]["style_id"], "Rubrik2")
        # Stycken utan pStyle får dokumentets standardstil
        self.assertEqual(self.by_id["paragraph_3"]["style_id"], "Normal")

    def test_doc_order_places_table_between_paragraphs(self):
        ordered = sorted(
            (e for e in self.structure["elements"] if e["doc_order"] is not None),
            key=lambda e: e["doc_order"],
        )
        self.assertEqual([e["element_id"] for e in ordered], [
            "paragraph_1", "paragraph_2", "paragraph_3", "paragraph_4", "paragraph_5",
            "table_1_cell_1_1_p1", "table_1_cell_1_2_p1",
            "paragraph_6", "paragraph_7",
        ])
        self.assertEqual([e["doc_order"] for e in ordered], list(range(1, 10)))


class UsageTrackerTests(unittest.TestCase):
    def test_reads_sdk_objects_dicts_and_missing_usage(self):
        tracker = UsageTracker()
        tracker.record("local", "m", SimpleNamespace(
            prompt_tokens=100, completion_tokens=40, total_tokens=140,
            completion_tokens_details=SimpleNamespace(reasoning_tokens=25),
            prompt_tokens_details=SimpleNamespace(cached_tokens=10),
        ))
        tracker.record("global", "m", {"input_tokens": 50, "output_tokens": 5})
        tracker.record("local", "m", None)
        tracker.record_failure("local", "m")

        local = tracker.totals("local")
        self.assertEqual((local.calls, local.failed_calls), (3, 1))
        self.assertEqual((local.prompt_tokens, local.completion_tokens), (100, 40))
        self.assertEqual((local.reasoning_tokens, local.cached_prompt_tokens), (25, 10))

        data = tracker.to_dict()
        self.assertEqual(data["total"]["prompt_tokens"], 150)
        self.assertEqual(data["by_phase"]["global"]["total_tokens"], 55)
        self.assertEqual(list(data["by_phase"]), ["local", "global"])


class _FakeCompletions:
    def __init__(self, content, usage):
        self.content = content
        self.usage = usage

    def create(self, **kwargs):
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=self.content))],
            usage=self.usage,
        )


class EndToEndRunSummaryTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"e2e-test-{id(self)}")
        self.source = self.root / "rapport.docx"

        document = Document()
        document.add_paragraph("Inledning", style="Heading 1")
        document.add_paragraph("Vi har genomfört en granskning av ärendet.")
        document.save(self.source)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_run_summary_contains_usage_counts_and_lix(self):
        content = json.dumps([{
            "type": "paragraph",
            "element_id": "paragraph_2",
            "old": "genomfört en granskning av",
            "new": "granskat",
            "motivation": "Verb i stället för substantiv.",
        }], ensure_ascii=False)
        usage = SimpleNamespace(prompt_tokens=1200, completion_tokens=80, total_tokens=1280)
        fake_client = SimpleNamespace(chat=SimpleNamespace(completions=_FakeCompletions(content, usage)))

        with mock.patch.object(suggestor_module.openai, "OpenAI", return_value=fake_client):
            improver = JBGLanguageImprover(
                input_path=str(self.source),
                api_key="unused",
                model="test-model",
                prompt_policy="policy",
                temperature=1,
                include_motivations=True,
                logger=self.logger,
                docx_mode="tracked",
            )
            improver.run(output_path=str(self.root / "out.docx"))

        summary = json.loads(Path(improver.run_summary_json).read_text(encoding="utf-8"))
        self.assertTrue(summary["succeeded"])
        self.assertEqual(summary["model"], "test-model")
        self.assertEqual(summary["usage"]["total"]["calls"], 1)
        self.assertEqual(summary["usage"]["by_phase"]["local"]["prompt_tokens"], 1200)
        self.assertEqual(summary["usage"]["by_phase"]["local"]["completion_tokens"], 80)

        counts = summary["local_suggestions"]
        self.assertEqual((counts["raw"], counts["accepted"], counts["applied"]), (1, 1, 1))

        readability = summary["readability"]
        self.assertEqual(readability["before"]["words"], 7)
        self.assertEqual(readability["after"]["words"], 4)
        self.assertEqual(readability["sections"][0]["heading"], "Inledning")

    def test_run_summary_is_saved_when_the_run_fails(self):
        with mock.patch.object(
            JBGLanguageImprover, "_build_change_plans", side_effect=RuntimeError("boom")
        ), mock.patch.object(
            suggestor_module.openai, "OpenAI",
            return_value=SimpleNamespace(chat=SimpleNamespace(
                completions=_FakeCompletions("[]", None))),
        ):
            improver = JBGLanguageImprover(
                input_path=str(self.source), api_key="unused", model="m",
                prompt_policy="p", temperature=1, include_motivations=False,
                logger=self.logger,
            )
            with self.assertRaises(RuntimeError):
                improver.run()

        summary = json.loads(Path(improver.run_summary_json).read_text(encoding="utf-8"))
        self.assertFalse(summary["succeeded"])
        self.assertEqual(summary["error"], "boom")


if __name__ == "__main__":
    unittest.main()
