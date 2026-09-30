import copy
import json
import logging
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from app.src import JBGLangImprovSuggestorAI as suggestor_module
from app.src.JBGAboutSectionRenderer import (
    ABOUT_BOOKMARK_NAME,
    ABOUT_HEADING,
    AboutSectionRenderer,
    build_about_section_blocks,
)
from app.src.JBGDocumentStructureExtractor import DocumentStructureExtractor
from app.src.JBGDocxPackage import DocxPackage
from app.src.JBGLanguageImprover import JBGLanguageImprover


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"
NBSP = "\u00a0"


def _quiet_logger(name):
    logger = logging.getLogger(name)
    logger.handlers.clear()
    logger.addHandler(logging.NullHandler())
    logger.propagate = False
    return logger


# Summering motsvarande testkörningen 93dbc428
REAL_RUN_SUMMARY = {
    "model": "gpt-5.2",
    "docx_mode": "tracked",
    "include_motivations": True,
    "prompt_customized": False,
    "started_at": "2026-09-30T06:45:34+00:00",
    "local_suggestions": {"raw": 134, "validated": 91, "accepted": 91, "planned": 91,
                          "applied": 91, "failed": 0, "comments_applied": 91},
    "usage": {"total": {"calls": 11, "failed_calls": 0, "prompt_tokens": 99322,
                        "completion_tokens": 30939, "reasoning_tokens": 0}},
    "readability": {"before": {"lix": 51.7, "band": "Svår"},
                    "after": {"lix": 51.6, "band": "Svår"}},
}


# ============================================================================
# Hjälpare för fixturer och simulering av Words "avvisa/godkänn alla"
# ============================================================================

def _style(document, name, outline_level=None, num_id=None):
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    ppr = style.element.get_or_add_pPr()
    if num_id is not None:
        num_pr = OxmlElement("w:numPr")
        num = OxmlElement("w:numId")
        num.set(qn("w:val"), str(num_id))
        num_pr.append(num)
        ppr.append(num_pr)
    if outline_level is not None:
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), str(outline_level))
        ppr.append(outline)
    return style


def build_iaf_like_document(path, with_sources=True, existing_revision_id=None):
    """Mall-lik fixtur: numrerade kapitel, onumrerade bilagerubriker, baksida."""
    document = Document()
    _style(document, "IAF Sammanfattning", outline_level=0)
    _style(document, "IAF Rubrik 1 numrerad - Kapitel", outline_level=0, num_id=5)
    _style(document, "IAF Bilagerubrik 1", outline_level=0)
    _style(document, "IAF Tabellrubrik", outline_level=0)
    _style(document, "IAF nRubrik 2", outline_level=0)

    document.add_paragraph("Sammanfattning", style="IAF Sammanfattning")
    document.add_paragraph("Myndigheten brister i arbetet.")
    document.add_paragraph("Slutsatser", style="IAF Rubrik 1 numrerad - Kapitel")
    body_text = document.add_paragraph("Vi har genomfört en granskning av ärendet.")
    document.add_paragraph("Tabell 1: Antal ärenden", style="IAF Tabellrubrik")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Totalt 1 200 ärenden."
    if with_sources:
        document.add_paragraph("Källor", style="IAF Bilagerubrik 1")
        document.add_paragraph("Förordning (2000:634).")
    document.add_paragraph("Vi är IAF", style="IAF nRubrik 2")
    back = document.add_paragraph("Telefon 010-123 45 67")
    back.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if existing_revision_id is not None:
        # Ett redan spårat ord i originalet, för kontroll av unika id:n
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), str(existing_revision_id))
        ins.set(qn("w:author"), "Någon")
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = " (spårad)"
        run.append(t)
        ins.append(run)
        body_text._p.append(ins)

    document.save(path)


def _body(docx_path):
    with ZipFile(docx_path) as z:
        return etree.fromstring(z.read("word/document.xml")).find(f"{W}body")


def _normalized(element):
    clone = copy.deepcopy(element)
    for node in clone.iter():
        if node.tail is not None and not node.tail.strip():
            node.tail = None
        if node.tag != f"{W}t" and node.text is not None and not node.text.strip():
            node.text = None
    return etree.tostring(clone, method="c14n")


def reject_all_insertions(body):
    """Som Words 'Avvisa alla infogningar': infogad text och styckemarkeringar försvinner."""
    for ins in list(body.iter(f"{W}ins")):
        if ins.getparent().tag != f"{W}rPr":
            ins.getparent().remove(ins)
    for paragraph in list(body.findall(f"{W}p")):
        if paragraph.find(f"{W}pPr/{W}rPr/{W}ins") is None:
            continue
        following = paragraph.getnext()
        content = [child for child in paragraph if child.tag != f"{W}pPr"]
        ppr = following.find(f"{W}pPr")
        position = 0 if ppr is None else list(following).index(ppr) + 1
        for child in reversed(content):
            following.insert(position, child)
        body.remove(paragraph)
    return body


def accept_all_insertions(body):
    for ins in list(body.iter(f"{W}ins")):
        parent = ins.getparent()
        if parent.tag == f"{W}rPr":
            parent.remove(ins)
            continue
        position = list(parent).index(ins)
        for child in reversed(list(ins)):
            parent.insert(position, child)
        parent.remove(ins)
    for rpr in list(body.iter(f"{W}rPr")):
        if rpr.getparent().tag == f"{W}pPr" and len(rpr) == 0:
            rpr.getparent().remove(rpr)
    return body


def _write_body(source_docx, target_docx, body):
    with ZipFile(source_docx) as zin, ZipFile(target_docx, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                root = etree.fromstring(data)
                root.replace(root.find(f"{W}body"), body)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)


# ============================================================================
# Innehåll
# ============================================================================

class AboutSectionContentTests(unittest.TestCase):
    def texts(self, summary, tracked=True):
        return [block.text for block in build_about_section_blocks(summary, tracked=tracked)]

    def test_real_run_values_are_formatted_in_swedish(self):
        texts = self.texts(REAL_RUN_SUMMARY)
        self.assertEqual(texts[0], ABOUT_HEADING)
        for expected in (
            "Datum: 30 september 2026",
            "Språkmodell: gpt-5.2",
            "Visning av förslagen: Spåra ändringar med kommentarer",
            "Promptinstruktion: standard",
            "Anrop till språkmodellen: 11",
            f"Tokens skickade: 99{NBSP}322",
            f"Tokens mottagna: 30{NBSP}939",
            "Förslag från språkmodellen: 134",
            "Förslag som klarade kontrollerna: 91",
            "Förslag som förts in i dokumentet: 91",
            "Läsbarhet (LIX) före: 51,7 (svår)",
            "Läsbarhet (LIX) om alla förslag godtas: 51,6 (svår)",
        ):
            self.assertIn(expected, texts)
        self.assertIn("avvisa ändringen", texts[1])
        self.assertFalse(any("resonemang" in t for t in texts))
        self.assertFalse(any("inte kunde föras in" in t for t in texts))

    def test_cached_tokens_are_shown_when_present(self):
        # Testkörningen ba144f3d: 56 704 av 60 817 tokens kom från cache
        summary = json.loads(json.dumps(REAL_RUN_SUMMARY))
        summary["usage"]["total"].update(prompt_tokens=60817, cached_prompt_tokens=56704)
        texts = self.texts(summary)
        self.assertIn(
            f"Tokens skickade: 60{NBSP}817, varav 56{NBSP}704 återanvända (lägre kostnad)", texts
        )

    def test_labels_are_bold_and_values_are_not(self):
        blocks = build_about_section_blocks(REAL_RUN_SUMMARY, tracked=True)
        model = next(b for b in blocks if b.text.startswith("Språkmodell"))
        self.assertEqual(model.runs, [("Språkmodell: ", True), ("gpt-5.2", False)])

    def test_optional_lines(self):
        summary = json.loads(json.dumps(REAL_RUN_SUMMARY))
        summary["docx_mode"] = "simple"
        summary["prompt_customized"] = None
        summary["readability"] = None
        summary["usage"]["total"].update(reasoning_tokens=1200, failed_calls=1)
        summary["local_suggestions"]["failed"] = 2
        texts = self.texts(summary, tracked=False)

        self.assertIn("Visning av förslagen: Enkel färgmarkering", texts)
        self.assertIn(f"Tokens mottagna: 30{NBSP}939, varav 1{NBSP}200 för resonemang", texts)
        self.assertIn("Anrop till språkmodellen: 11, varav 1 misslyckades", texts)
        self.assertIn(f"Tokens skickade: 99{NBSP}322", texts)   # inga cachade tokens
        self.assertIn("Förslag som inte kunde föras in: 2", texts)
        self.assertFalse(any(t.startswith("Promptinstruktion") for t in texts))
        self.assertFalse(any("LIX" in t for t in texts))
        self.assertIn("ta bort avsnittet när granskningen är klar", texts[1])


# ============================================================================
# Rendering
# ============================================================================

class AboutSectionRenderingTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"about-test-{id(self)}")
        self.source = self.root / "rapport.docx"

    def tearDown(self):
        self.temp_dir.cleanup()

    def render(self, source, output, tracked, summary=REAL_RUN_SUMMARY):
        blocks = build_about_section_blocks(summary, tracked=tracked)
        with DocxPackage(str(source), self.logger) as pkg:
            result = AboutSectionRenderer(pkg, self.logger).apply(blocks, tracked=tracked)
            pkg.save(str(output))
        return blocks, result

    def test_plain_section_is_last_with_unnumbered_appendix_style(self):
        build_iaf_like_document(self.source)
        output = self.root / "simple.docx"
        blocks, result = self.render(self.source, output, tracked=False)

        self.assertTrue(result.applied)
        self.assertEqual(result.heading_style_id, "IAFBilagerubrik1")
        body = _body(output)
        children = list(body)
        self.assertEqual(children[-1].tag, f"{W}sectPr")

        paragraphs = body.findall(f"{W}p")
        section = paragraphs[-len(blocks):]
        self.assertEqual(paragraphs[-len(blocks) - 1].findtext(f".//{W}t"), "Telefon 010-123 45 67")
        heading_ppr = section[0].find(f"{W}pPr")
        self.assertEqual(heading_ppr.find(f"{W}pStyle").get(f"{W}val"), "IAFBilagerubrik1")
        self.assertIsNotNone(heading_ppr.find(f"{W}pageBreakBefore"))
        self.assertEqual(list(body.iter(f"{W}ins")), [])
        self.assertEqual(
            [b.get(f"{W}name") for b in body.iter(f"{W}bookmarkStart")], [ABOUT_BOOKMARK_NAME]
        )
        Document(str(output))  # öppnas utan fel

    def test_first_unnumbered_heading_style_is_used_without_sources_heading(self):
        build_iaf_like_document(self.source, with_sources=False)
        _, result = self.render(self.source, self.root / "out.docx", tracked=False)
        self.assertEqual(result.heading_style_id, "IAFSammanfattning")

    def test_numbered_style_is_used_only_with_numbering_blocked(self):
        document = Document()
        _style(document, "Kapitel", outline_level=0, num_id=3)
        _style(document, "Numrerad bas", num_id=4)
        document.styles["Heading 1"].base_style = document.styles["Numrerad bas"]
        document.add_paragraph("Inledning", style="Kapitel")
        document.add_paragraph("Text.")
        document.save(self.source)

        output = self.root / "out.docx"
        blocks, result = self.render(self.source, output, tracked=False)
        self.assertEqual(result.heading_style_id, "Kapitel")
        heading = _body(output).findall(f"{W}p")[-len(blocks)]
        self.assertEqual(heading.find(f"{W}pPr/{W}numPr/{W}numId").get(f"{W}val"), "0")

    def test_bold_fallback_when_the_document_has_no_heading_styles(self):
        document = Document()
        styles_root = document.styles.element
        for style in list(styles_root):
            if style.get(qn("w:styleId")) in {"Heading1"}:
                styles_root.remove(style)
        document.add_paragraph("Bara brödtext.")
        document.save(self.source)

        output = self.root / "out.docx"
        blocks, result = self.render(self.source, output, tracked=False)
        self.assertIsNone(result.heading_style_id)
        heading = _body(output).findall(f"{W}p")[-len(blocks)]
        self.assertIsNone(heading.find(f"{W}pPr/{W}pStyle"))
        self.assertIsNotNone(heading.find(f"{W}r/{W}rPr/{W}b"))

    def test_tracked_section_is_one_insertion_with_unique_ids(self):
        build_iaf_like_document(self.source, existing_revision_id=41)
        output = self.root / "tracked.docx"
        blocks, _ = self.render(self.source, output, tracked=True)
        body = _body(output)

        # Alla nya körningar ligger i w:ins med id över originalets största (41)
        section_runs = [
            r for r in body.iter(f"{W}r")
            if (r.findtext(f"{W}t") or "") in {run for b in blocks for run, _ in b.runs}
        ]
        self.assertTrue(section_runs)
        self.assertTrue(all(r.getparent().tag == f"{W}ins" for r in section_runs))
        new_ids = [
            int(i.get(f"{W}id")) for i in body.iter(f"{W}ins") if i.get(f"{W}author") != "Någon"
        ]
        self.assertTrue(all(i > 41 for i in new_ids))
        self.assertEqual(len(new_ids), len(set(new_ids)))

        # Baksidans sista stycke har fått infogad styckemarkering
        back = next(p for p in body.findall(f"{W}p") if p.findtext(f".//{W}t") == "Telefon 010-123 45 67")
        self.assertIsNotNone(back.find(f"{W}pPr/{W}rPr/{W}ins"))
        # Bokmärket ligger inne i infogningen
        start = next(body.iter(f"{W}bookmarkStart"))
        self.assertEqual(start.getparent().tag, f"{W}ins")
        Document(str(output))

    def test_rejecting_the_insertion_restores_the_original_exactly(self):
        build_iaf_like_document(self.source)
        output = self.root / "tracked.docx"
        self.render(self.source, output, tracked=True)

        rejected = reject_all_insertions(_body(output))
        self.assertEqual(_normalized(rejected), _normalized(_body(self.source)))

    def test_accepting_keeps_the_section_and_a_rerun_neither_reviews_nor_duplicates_it(self):
        build_iaf_like_document(self.source)
        tracked = self.root / "tracked.docx"
        self.render(self.source, tracked, tracked=True)

        accepted = self.root / "accepted.docx"
        _write_body(tracked, accepted, accept_all_insertions(_body(tracked)))
        Document(str(accepted))

        original_ids = [
            e["element_id"] for e in DocumentStructureExtractor(str(self.source), self.logger).extract()["elements"]
        ]
        structure = DocumentStructureExtractor(str(accepted), self.logger).extract()
        texts = [e["text"] for e in structure["elements"]]
        self.assertNotIn(ABOUT_HEADING, texts)
        self.assertFalse(any(t.startswith("Språkmodell") for t in texts))
        self.assertIn("Telefon 010-123 45 67", texts)  # baksidan granskas fortfarande
        # Originalets element behåller id och ordning; tillkommet är bara det
        # tomma avslutande stycket efter avsnittet.
        extracted_ids = [e["element_id"] for e in structure["elements"]]
        self.assertEqual([i for i in extracted_ids if i in original_ids], original_ids)
        extra = [e for e in structure["elements"] if e["element_id"] not in original_ids]
        self.assertTrue(all(e["empty"] for e in extra), extra)

        rerun = self.root / "rerun.docx"
        _, result = self.render(accepted, rerun, tracked=False)
        self.assertTrue(result.replaced_existing)
        body = _body(rerun)
        headings = [p for p in body.findall(f"{W}p") if p.findtext(f".//{W}t") == ABOUT_HEADING]
        self.assertEqual(len(headings), 1)
        self.assertEqual(len([b for b in body.iter(f"{W}bookmarkStart")
                              if b.get(f"{W}name") == ABOUT_BOOKMARK_NAME]), 1)
        # Inga ackumulerade tomma stycken före avsnittet
        previous = headings[0].getprevious()
        self.assertEqual(previous.findtext(f".//{W}t"), "Telefon 010-123 45 67")


# ============================================================================
# Samma tre regler för alla typer av dokument
# ============================================================================

def _doc_plain(path):
    document = Document()
    document.add_paragraph("Ett dokument utan rubriker.")
    document.add_paragraph("Andra stycket.")
    document.save(path)


def _doc_numbered_heading_via_style(path):
    document = Document()
    _style(document, "Numrerad bas", num_id=1)
    document.styles["Heading 1"].base_style = document.styles["Numrerad bas"]
    document.add_paragraph("Inledning", style="Heading 1")
    document.add_paragraph("Text.")
    document.save(path)


def _doc_numbered_heading_via_numbering_part(path):
    """Numreringen kopplas från numbering.xml (w:lvl/w:pStyle), inte i stilen."""
    document = Document()
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "90")
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "decimal"), ("w:pStyle", "Heading1"),
                       ("w:lvlText", "%1")):
        child = OxmlElement(tag)
        child.set(qn("w:val"), value)
        lvl.append(child)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    document.add_paragraph("Bakgrund", style="Heading 1")
    document.add_paragraph("Text.")
    document.save(path)


def _doc_english_references(path):
    document = Document()
    document.add_paragraph("Introduction", style="Heading 1")
    document.add_paragraph("Some text.")
    document.add_paragraph("References", style="Heading 1")
    document.add_paragraph("A book.")
    document.save(path)


def _doc_ends_with_table(path):
    document = Document()
    document.add_paragraph("Text före tabellen.")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "Sista cellen"
    document.save(path)


def _doc_ends_with_content_control(path):
    document = Document()
    document.add_paragraph("Text före innehållskontrollen.")
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Innehåll i kontroll"
    run.append(t)
    paragraph.append(run)
    content.append(paragraph)
    sdt.append(content)
    document.element.body.sectPr.addprevious(sdt)
    document.save(path)


def _doc_ends_with_section_break(path):
    document = Document()
    document.add_paragraph("Första avsnittet.")
    document.add_section()          # sectPr i sista styckets pPr
    document.save(path)


def _doc_empty(path):
    Document().save(path)


def _doc_last_paragraph_without_ppr(path):
    document = Document()
    document.add_paragraph("Enkelt sista stycke utan styckeformat.")
    document.save(path)


DOCUMENT_SHAPES = {
    "iaf_like": build_iaf_like_document,
    "plain": _doc_plain,
    "numbered_heading_via_style": _doc_numbered_heading_via_style,
    "numbered_heading_via_numbering_part": _doc_numbered_heading_via_numbering_part,
    "english_references": _doc_english_references,
    "ends_with_table": _doc_ends_with_table,
    "ends_with_content_control": _doc_ends_with_content_control,
    "ends_with_section_break": _doc_ends_with_section_break,
    "empty": _doc_empty,
    "last_paragraph_without_ppr": _doc_last_paragraph_without_ppr,
}

# Dokument som inte slutar med ett vanligt stycke får ett tomt avslutande
# stycke kvar efter avvisning (Word kräver ett stycke sist i dokumentet).
SHAPES_WITH_TRAILING_EMPTY_PARAGRAPH = {
    "ends_with_table", "ends_with_content_control", "ends_with_section_break", "empty",
}


class AboutSectionRulesForAllDocumentsTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"about-rules-{id(self)}")

    def tearDown(self):
        self.temp_dir.cleanup()

    def _render(self, shape, tracked):
        source = self.root / f"{shape}.docx"
        output = self.root / f"{shape}_{'tracked' if tracked else 'simple'}.docx"
        DOCUMENT_SHAPES[shape](str(source))
        blocks = build_about_section_blocks(REAL_RUN_SUMMARY, tracked=tracked)
        with DocxPackage(str(source), self.logger) as pkg:
            result = AboutSectionRenderer(pkg, self.logger).apply(blocks, tracked=tracked)
            pkg.save(str(output))
        return source, output, blocks, result

    def _heading(self, body):
        return next(p for p in body.findall(f"{W}p")
                    if "".join(t.text or "" for t in p.iter(f"{W}t")) == ABOUT_HEADING)

    def assert_last(self, body, blocks):
        """Regel 1: efter allt befintligt innehåll, före body-sectPr, på ny sida."""
        children = [c for c in body if c.tag != f"{W}sectPr"]
        heading = self._heading(body)
        section_start = children.index(heading)
        tail = children[section_start:]
        # Avsnittets stycken plus högst ett avslutande tomt stycke
        self.assertIn(len(tail), (len(blocks), len(blocks) + 1))
        self.assertTrue(all(c.tag == f"{W}p" for c in tail))
        self.assertEqual(list(body)[-1].tag, f"{W}sectPr")
        self.assertIsNotNone(heading.find(f"{W}pPr/{W}pageBreakBefore"))

    def assert_unnumbered(self, heading):
        """Regel 3: direkt numreringsspärr oavsett stil."""
        num_id = heading.find(f"{W}pPr/{W}numPr/{W}numId")
        self.assertIsNotNone(num_id)
        self.assertEqual(num_id.get(f"{W}val"), "0")

    def test_rules_hold_for_every_document_shape(self):
        for shape in DOCUMENT_SHAPES:
            for tracked in (False, True):
                with self.subTest(shape=shape, tracked=tracked):
                    source, output, blocks, result = self._render(shape, tracked)
                    self.assertTrue(result.applied)
                    Document(str(output))  # giltigt för python-docx
                    body = _body(output)

                    self.assert_last(body, blocks)
                    self.assert_unnumbered(self._heading(body))

                    section_texts = {text for b in blocks for text, _ in b.runs}
                    section_runs = [r for r in body.iter(f"{W}r")
                                    if (r.findtext(f"{W}t") or "") in section_texts]
                    self.assertEqual(len(section_runs), sum(len(b.runs) for b in blocks))
                    if tracked:
                        # Regel 2: allt i avsnittet är spårat
                        self.assertTrue(all(r.getparent().tag == f"{W}ins" for r in section_runs))
                    else:
                        self.assertEqual(list(body.iter(f"{W}ins")), [])

    def test_rejecting_restores_every_document_shape(self):
        for shape in DOCUMENT_SHAPES:
            with self.subTest(shape=shape):
                source, output, _, _ = self._render(shape, tracked=True)
                rejected = reject_all_insertions(_body(output))
                expected = _body(source)
                if shape in SHAPES_WITH_TRAILING_EMPTY_PARAGRAPH:
                    expected.find(f"{W}sectPr").addprevious(etree.Element(f"{W}p"))
                self.assertEqual(_normalized(rejected), _normalized(expected))
                self.assertEqual(
                    [b for b in rejected.iter(f"{W}bookmarkStart")
                     if b.get(f"{W}name") == ABOUT_BOOKMARK_NAME], []
                )

    def test_heading_style_choice_per_shape(self):
        expected = {
            "iaf_like": "IAFBilagerubrik1",                        # onumrerad, som Källor
            "plain": "Heading1",                                   # inbyggd, onumrerad
            "numbered_heading_via_style": "Heading1",              # numrerad -> spärras
            "numbered_heading_via_numbering_part": "Heading1",     # numrerad -> spärras
            "english_references": "Heading1",
        }
        for shape, style_id in expected.items():
            with self.subTest(shape=shape):
                _, _, _, result = self._render(shape, tracked=False)
                self.assertEqual(result.heading_style_id, style_id)

    def test_unnumbered_style_is_preferred_over_style_numbered_from_numbering_part(self):
        # Heading 1 kommer först men numreras via numbering.xml; den onumrerade
        # egna rubrikstilen längre ned ska väljas.
        source = self.root / "prefer.docx"
        _doc_numbered_heading_via_numbering_part(str(source))
        document = Document(str(source))
        _style(document, "Bilagerubrik", outline_level=0)
        document.add_paragraph("Sammanfattning av bilagan", style="Bilagerubrik")
        document.add_paragraph("Text.")
        document.save(str(source))

        with DocxPackage(str(source), self.logger) as pkg:
            result = AboutSectionRenderer(pkg, self.logger).apply(
                build_about_section_blocks(REAL_RUN_SUMMARY, tracked=False), tracked=False
            )
        self.assertEqual(result.heading_style_id, "Bilagerubrik")

    def test_numbering_linked_from_numbering_part_is_detected(self):
        source = self.root / "linked.docx"
        _doc_numbered_heading_via_numbering_part(str(source))
        with DocxPackage(str(source), self.logger) as pkg:
            renderer = AboutSectionRenderer(pkg, self.logger)
            linked = renderer._styles_numbered_by_numbering_part()
            styles = renderer._style_index()
            self.assertIn("Heading1", linked)
            self.assertTrue(renderer._resolved(styles, "Heading1", linked)["numbered"])


# ============================================================================
# Hela kedjan
# ============================================================================

class _FakeCompletions:
    def create(self, **kwargs):
        content = json.dumps([{
            "type": "paragraph", "element_id": "paragraph_4",
            "old": "genomfört en granskning av", "new": "granskat",
            "motivation": "Verb i stället för substantiv.",
        }], ensure_ascii=False)
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))],
            usage=SimpleNamespace(prompt_tokens=1500, completion_tokens=90, total_tokens=1590),
        )


class AboutSectionEndToEndTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"about-e2e-{id(self)}")
        self.source = self.root / "rapport.docx"
        build_iaf_like_document(self.source)

    def tearDown(self):
        self.temp_dir.cleanup()

    def run_improver(self, include_about_section, docx_mode="tracked"):
        client = SimpleNamespace(chat=SimpleNamespace(completions=_FakeCompletions()))
        output = self.root / f"out_{include_about_section}_{docx_mode}.docx"
        with mock.patch.object(suggestor_module.openai, "OpenAI", return_value=client):
            improver = JBGLanguageImprover(
                input_path=str(self.source), api_key="unused", model="test-model",
                prompt_policy="policy", temperature=1, include_motivations=True,
                logger=self.logger, docx_mode=docx_mode,
                include_about_section=include_about_section, prompt_customized=True,
            )
            improver.run(output_path=str(output))
        texts = [t.text or "" for t in _body(output).iter(f"{W}t")]
        return improver, texts

    def test_section_is_added_with_the_runs_own_numbers(self):
        for mode in ("tracked", "simple"):
            with self.subTest(mode=mode):
                improver, texts = self.run_improver(True, mode)
                self.assertIn(ABOUT_HEADING, texts)
                self.assertIn("test-model", texts)
                self.assertIn(f"1{NBSP}500", texts)             # tokens skickade
                self.assertIn("anpassad för den här granskningen", texts)
                about = improver.run_summary.about_section
                self.assertTrue(about["applied"], about)
                self.assertEqual(about["heading_style_id"], "IAFBilagerubrik1")
                self.assertEqual(improver.run_summary.local_suggestions.applied, 1)

    def test_lix_can_be_turned_off(self):
        client = SimpleNamespace(chat=SimpleNamespace(completions=_FakeCompletions()))
        output = self.root / "no_lix.docx"
        with mock.patch.object(suggestor_module.openai, "OpenAI", return_value=client):
            improver = JBGLanguageImprover(
                input_path=str(self.source), api_key="unused", model="test-model",
                prompt_policy="policy", temperature=1, include_motivations=True,
                logger=self.logger, docx_mode="simple", compute_readability=False,
            )
            improver.run(output_path=str(output))
        texts = [t.text or "" for t in _body(output).iter(f"{W}t")]
        self.assertIn(ABOUT_HEADING, texts)
        self.assertFalse(any("LIX" in t for t in texts))
        self.assertIsNone(improver.run_summary.readability)
        self.assertFalse(improver.run_summary.compute_readability)

    def test_lix_is_computed_without_the_section(self):
        improver, texts = self.run_improver(False)
        self.assertNotIn(ABOUT_HEADING, texts)
        self.assertIsNotNone(improver.run_summary.readability)

    def test_section_is_not_added_when_unchecked(self):
        improver, texts = self.run_improver(False)
        self.assertNotIn(ABOUT_HEADING, texts)
        self.assertIsNone(improver.run_summary.about_section)


if __name__ == "__main__":
    unittest.main()
