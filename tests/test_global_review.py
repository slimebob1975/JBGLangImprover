import json
import logging
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from zipfile import ZipFile

from docx import Document
from lxml import etree

from app.src import JBGGlobalAnalyzerAI as analyzer_module
from app.src import JBGLangImprovSuggestorAI as suggestor_module
from app.src.JBGDocumentStructureExtractor import DocumentStructureExtractor
from app.src.JBGDocxPackage import DocxPackage
from app.src.JBGGlobalAnalyzerAI import (
    GlobalFinding,
    GlobalReviewResult,
    JBGGlobalAnalyzerAI,
    build_outline,
    load_global_policy,
    split_outline,
)
from app.src.JBGGlobalFindingsRenderer import GlobalFindingsRenderer
from app.src.JBGLanguageImprover import JBGLanguageImprover
from app.src.JBGUsageTracker import UsageTracker
from tests.test_about_section import about_tables


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


def _quiet_logger(name):
    logger = logging.getLogger(name)
    logger.handlers.clear()
    logger.addHandler(logging.NullHandler())
    logger.propagate = False
    return logger


def _el(eid, etype, text, level=None, order=None, **extra):
    element = {"type": etype, "element_id": eid, "text": text, "heading_level": level,
               "doc_order": order, "style_id": "Normal", "style_name": "Normal"}
    element.update(extra)
    return element


STRUCTURE = {"type": "docx", "elements": [
    _el("paragraph_1", "paragraph", "Sammanfattning", 1, 1),
    _el("paragraph_2", "paragraph", "IAF har granskat hur a-kassorna kontrollerar tidrapporter.", None, 2),
    _el("paragraph_3", "paragraph", "Inledning", 1, 3),
    _el("paragraph_4", "paragraph", "Granskningen omfattar alla a-kassor under 2024.", None, 4),
    _el("table_1_cell_1_1_p1", "table_cell", "Tabellen visar  antalet kontroller.", None, 5),
    _el("paragraph_5", "paragraph", "Resultat", 1, 6),
    _el("paragraph_6", "paragraph", "Granskningen omfattar alla a-kassor under 2024, vilket nämndes ovan.", None, 7),
    _el("paragraph_7", "paragraph", "", None, 8),
    _el("paragraph_8", "paragraph", "Andelen tidrapporter med återkrav var 2,2 procent under 2021–2023.", None, 9),
    _el("paragraph_9", "paragraph", "Under 2021–2023 ledde 3,1 procent av tidrapporterna till återkrav.", None, 10),
    _el("textbox_2_p1", "textbox", "40–63 % ha diagranm%", None, 11,
        container_path="/document/body/paragraph[9]/textbox[1]/paragraph[1]"),
    _el("footnote_1", "footnote", "SFS 1997:238.", footnote_id="2"),
    _el("header_1_p1", "header", "Rapport 2025:3"),
]}


class OutlineTests(unittest.TestCase):
    def test_outline_is_body_text_in_reading_order(self):
        outline = build_outline(STRUCTURE)
        self.assertEqual([r["id"] for r in outline], [
            "paragraph_1", "paragraph_2", "paragraph_3", "paragraph_4",
            "table_1_cell_1_1_p1", "paragraph_5", "paragraph_6", "paragraph_8", "paragraph_9",
            "textbox_2_p1",
        ])
        self.assertEqual(outline[0], {"id": "paragraph_1", "h": 1, "t": "Sammanfattning"})
        self.assertNotIn("h", outline[1])
        self.assertEqual(outline[4]["t"], "Tabellen visar antalet kontroller.")  # blanksteg normaliserade

    def test_small_documents_are_one_part(self):
        self.assertEqual(len(split_outline(STRUCTURE, build_outline(STRUCTURE))), 1)

    def test_large_documents_split_at_top_level_headings(self):
        # Avsnitten är 142, 199 och 340 tecken: med gränsen 400 ryms alla
        # avsnitt var för sig, så varje del börjar med en avsnittsrubrik.
        outline = build_outline(STRUCTURE)
        parts = split_outline(STRUCTURE, outline, max_chars=400)
        self.assertEqual([part[0]["id"] for part in parts], ["paragraph_1", "paragraph_5"])
        self.assertEqual([r["id"] for part in parts for r in part], [r["id"] for r in outline])

    def test_a_section_larger_than_the_limit_is_split_between_elements(self):
        outline = build_outline(STRUCTURE)
        parts = split_outline(STRUCTURE, outline, max_chars=300)
        self.assertEqual([r["id"] for part in parts for r in part], [r["id"] for r in outline])
        self.assertIn("paragraph_9", [part[0]["id"] for part in parts])

    def test_policy_loads_without_markers_and_defines_the_output_format(self):
        policy = load_global_policy()
        self.assertNotIn("<!--", policy)
        for required in ('"findings"', '"category"', '"element_ids"', '"quote"', "repetition"):
            self.assertIn(required, policy)
        # Upprepning mellan detaljnivåer är avsiktlig, åt båda hållen
        self.assertIn("olika nivåer är avsiktligt", policy)
        self.assertIn("oavsett vilket av ställena som kommer först", policy)
        self.assertIn("utelämningstecken", policy)
        # Inkonsekvenser: kategori, krav på related_quote och kända undantag
        self.assertIn("inconsistency", policy)
        self.assertIn('"related_quote"', policy)
        self.assertIn("avrundningar av samma värde", policy)
        self.assertIn("numreringen inte syns i texten", policy)
        # Troliga fel: ett ställe räcker, stavfel hör till den lokala granskningen
        self.assertIn("### error", policy)
        self.assertIn('Kategorierna "error", "disposition" och "heading" kan ha ett enda id', policy)
        # Disposition och rubriker: respektfull ton, skyddade standardavsnitt,
        # övergångar bara som stöd, gräns mot granskningen av belägg
        self.assertIn("### disposition", policy)
        self.assertIn("### heading", policy)
        self.assertIn("respektfulla förslag", policy)
        self.assertIn("ska behålla sina rubriker och sin plats", policy)
        self.assertIn("aldrig som en egen iakttagelse", policy)
        self.assertIn("Bedöm inte om ett påstående i rubriken är belagt", policy)
        self.assertIn('"proposed_order"', policy)
        self.assertIn("hanteras i den lokala granskningen", policy)


class ParseAndValidateTests(unittest.TestCase):
    def setUp(self):
        self.analyzer = JBGGlobalAnalyzerAI(
            api_key="unused", model="m", temperature=1,
            logger=_quiet_logger(f"global-validate-{id(self)}"), policy="p",
        )

    def test_parse_accepts_object_list_and_fenced_json(self):
        item = {"category": "repetition"}
        self.assertEqual(self.analyzer.parse_response(json.dumps({"findings": [item]})), [item])
        self.assertEqual(self.analyzer.parse_response(json.dumps([item])), [item])
        self.assertEqual(self.analyzer.parse_response("```json\n" + json.dumps({"findings": [item]}) + "\n```"), [item])
        self.assertEqual(self.analyzer.parse_response('Här är svaret: {"findings": []}'), [])
        with self.assertRaises(ValueError):
            self.analyzer.parse_response("inget json alls")

    def validate(self, *items):
        result = GlobalReviewResult()
        self.analyzer.validate(STRUCTURE, [(1, item) for item in items], result)
        return result

    def repetition(self, **overrides):
        item = {
            "category": "repetition",
            "element_ids": ["paragraph_6", "paragraph_4"],
            "quote": "Granskningen omfattar alla a-kassor under 2024, vilket nämndes ovan.",
            "description": "Samma avgränsning står redan i inledningen.",
            "proposal": "Stryk meningen.",
        }
        item.update(overrides)
        return item

    def test_valid_repetition_is_accepted(self):
        result = self.validate(self.repetition())
        self.assertEqual(len(result.findings), 1)
        finding = result.findings[0]
        self.assertEqual(finding.element_ids, ["paragraph_6", "paragraph_4"])
        self.assertEqual(finding.label, "Onödig upprepning")

    def test_quote_in_another_element_makes_that_element_the_anchor(self):
        result = self.validate(self.repetition(quote="Granskningen omfattar alla a-kassor under 2024."))
        self.assertEqual(result.findings[0].element_ids, ["paragraph_4", "paragraph_6"])

    def test_whitespace_differences_in_quote_are_tolerated(self):
        result = self.validate(self.repetition(
            element_ids=["table_1_cell_1_1_p1", "paragraph_2"], quote="Tabellen visar antalet"))
        self.assertEqual(len(result.findings), 1)

    def test_quotes_with_ellipsis_are_matched_fragment_by_fragment(self):
        # Testkörningen 6ea91865: tre riktiga citat avvisades för att de innehöll "..."
        text = ("Baserat på de förbättringar av arbetslöshetskassornas kontroller som vi redogör för i "
                "kapitel 4.4, anser IAF det rimligt att anta att andelen hittade felaktiga utbetalningar "
                "ligger närmare den övre än den nedre gränsen av de uppskattade intervallen, det vill säga "
                "omkring två tredjedelar av de felaktiga utbetalningarna.")
        contains = self.analyzer._contains
        self.assertTrue(contains(text, "... ligger närmare den övre än den nedre gränsen ... det vill säga "
                                       "omkring två tredjedelar av de felaktiga utbetalningarna."))
        self.assertTrue(contains(text, "Baserat på de förbättringar … omkring två tredjedelar"))
        # Testkörningen 3bb64fd6: citat som bara inleds med "..."
        self.assertTrue(contains(text, "...det vill säga omkring två tredjedelar av de felaktiga utbetalningarna."))
        # Delar i fel ordning, delar som saknas och bara korta delar godtas inte
        self.assertFalse(contains(text, "omkring två tredjedelar ... Baserat på de förbättringar"))
        self.assertFalse(contains(text, "ligger närmare den övre ... en mening som inte finns"))
        self.assertFalse(contains(text, "IAF ... det ... två"))

    def test_rejections_have_reasons(self):
        result = self.validate(
            self.repetition(category="structure_order"),
            self.repetition(element_ids=["paragraph_999"]),
            self.repetition(quote="Den här meningen finns inte."),
            self.repetition(element_ids=["paragraph_6"]),
            self.repetition(description=""),
            self.repetition(element_ids=["footnote_1", "paragraph_4"], quote="SFS 1997:238."),
            "inte ett objekt",
        )
        self.assertEqual(result.findings, [])
        self.assertEqual([r.reason for r in result.rejected], [
            "unsupported_category", "unknown_element_ids", "quote_not_found",
            "repetition_needs_two_locations", "missing_description", "quote_not_found",
            "not_an_object",
        ])

    def inconsistency(self, **overrides):
        item = {
            "category": "inconsistency",
            "element_ids": ["paragraph_9", "paragraph_8"],
            "quote": "ledde 3,1 procent av tidrapporterna till återkrav",
            "related_quote": "var 2,2 procent under 2021–2023",
            "description": "Andelen tidrapporter med återkrav anges olika för samma period.",
            "proposal": "Kontrollera siffran och använd samma värde på båda ställena.",
        }
        item.update(overrides)
        return item

    def test_valid_inconsistency_is_accepted_with_both_quotes(self):
        result = self.validate(self.inconsistency())
        self.assertEqual(len(result.findings), 1)
        finding = result.findings[0]
        self.assertEqual(finding.label, "Inkonsekvent påstående")
        self.assertEqual(finding.element_ids, ["paragraph_9", "paragraph_8"])
        self.assertEqual(finding.related_quote, "var 2,2 procent under 2021–2023")

    def test_related_quote_decides_which_place_comes_second(self):
        result = self.validate(self.inconsistency(element_ids=["paragraph_9", "paragraph_2", "paragraph_8"]))
        self.assertEqual(result.findings[0].element_ids, ["paragraph_9", "paragraph_8", "paragraph_2"])

    def test_inconsistency_requires_a_verified_related_quote(self):
        result = self.validate(
            self.inconsistency(related_quote=""),
            self.inconsistency(related_quote="4,0 procent"),
            self.inconsistency(element_ids=["paragraph_9"]),
            # Det motstridiga citatet måste finnas på ett annat ställe än ankaret
            self.inconsistency(related_quote="ledde 3,1 procent av tidrapporterna"),
        )
        self.assertEqual(result.findings, [])
        self.assertEqual([r.reason for r in result.rejected], [
            "missing_related_quote", "related_quote_not_found",
            "inconsistency_needs_two_locations", "related_quote_not_found",
        ])

    def test_optional_related_quote_that_does_not_match_is_dropped_for_repetitions(self):
        result = self.validate(self.repetition(related_quote="Den här meningen finns inte."))
        self.assertEqual(len(result.findings), 1)
        self.assertEqual(result.findings[0].related_quote, "")

    def test_error_in_a_single_place_is_accepted(self):
        # Testkörningen 7c5dfd4d: kvarlämnad arbetstext i en textruta
        result = self.validate({
            "category": "error",
            "element_ids": ["textbox_2_p1"],
            "quote": "40–63 % ha diagranm%",
            "description": "Textrutan innehåller kvarlämnad arbetstext.",
            "proposal": "Ta bort ”ha diagranm%”.",
        })
        self.assertEqual(len(result.findings), 1, [r.reason for r in result.rejected])
        self.assertEqual(result.findings[0].label, "Troligt fel")
        self.assertEqual(result.findings[0].element_ids, ["textbox_2_p1"])

    def test_error_quote_must_still_be_found(self):
        result = self.validate({
            "category": "error", "element_ids": ["textbox_2_p1"],
            "quote": "TODO: infoga siffra", "description": "Platshållare.",
        })
        self.assertEqual([r.reason for r in result.rejected], ["quote_not_found"])

    # ---------------- Disposition och rubriker ----------------

    def structure_with_caption(self):
        structure = json.loads(json.dumps(STRUCTURE))
        structure["elements"].append(_el(
            "paragraph_10", "paragraph", "Tabell 1: Antal kontroller", 1, 12,
            style_id="IAFTabellrubrik", style_name="IAF Tabellrubrik",
        ))
        return structure

    def validate_in(self, structure, *items):
        result = GlobalReviewResult()
        self.analyzer.validate(structure, [(1, item) for item in items], result)
        return result

    def disposition(self, **overrides):
        item = {
            "category": "disposition",
            "element_ids": ["paragraph_5", "paragraph_3"],
            "quote": "Resultat",
            "description": "Resultaten presenteras före den metod de bygger på.",
            "proposal": "Överväg att flytta avsnittet efter inledningen.",
            "proposed_order": ["Sammanfattning", "Inledning", "Resultat"],
        }
        item.update(overrides)
        return item

    def heading(self, **overrides):
        item = {
            "category": "heading",
            "element_ids": ["paragraph_3"],
            "quote": "Inledning",
            "description": "Avsnittet beskriver främst granskningens avgränsning.",
            "proposal": "Rubriken skulle kunna vara till exempel ”Granskningens avgränsning” (förslag).",
        }
        item.update(overrides)
        return item

    def test_disposition_and_heading_findings_are_accepted_on_headings(self):
        result = self.validate(self.disposition(), self.heading())
        self.assertEqual([f.label for f in result.findings],
                         ["Förslag om disposition", "Förslag om rubrik"])
        self.assertEqual(result.findings[0].proposed_order, ["Sammanfattning", "Inledning", "Resultat"])
        self.assertEqual(result.findings[1].element_ids, ["paragraph_3"])

    def test_comment_must_sit_on_a_heading_not_body_text_or_a_caption(self):
        result = self.validate_in(
            self.structure_with_caption(),
            self.heading(element_ids=["paragraph_4"], quote="Granskningen omfattar alla a-kassor"),
            self.heading(element_ids=["paragraph_10"], quote="Tabell 1: Antal kontroller"),
        )
        self.assertEqual(result.findings, [])
        self.assertEqual([r.reason for r in result.rejected],
                         ["anchor_not_a_heading", "anchor_not_a_heading"])

    def test_proposed_order_is_kept_only_when_all_headings_exist(self):
        result = self.validate(
            self.disposition(proposed_order=["sammanfattning ", "Resultat.", "Inledning"]),
            self.disposition(element_ids=["paragraph_3", "paragraph_1"], quote="Inledning",
                             proposed_order=["Inledning", "Metod", "Resultat"]),
        )
        # Skiftläge, blanksteg och avslutande punkt spelar ingen roll; texten
        # hämtas från dokumentets riktiga rubriker.
        self.assertEqual(result.findings[0].proposed_order, ["Sammanfattning", "Resultat", "Inledning"])
        # "Metod" finns inte som rubrik: ordningen tas bort, iakttagelsen behålls
        self.assertEqual(result.findings[1].proposed_order, [])
        self.assertEqual(len(result.findings), 2)

    def structure_with_fact_box(self):
        # Testkörningen b391e3ac: faktarutan "IAF:s tillsyn" är en tabell före förordet
        structure = json.loads(json.dumps(STRUCTURE))
        structure["elements"] += [
            _el("table_9_cell_1_1_p1", "table_cell", "IAF:s tillsyn", 1, 0,
                style_name="IAF Rubrik 1 - ej i innehållsförteckningen"),
            _el("table_9_cell_1_1_p2", "table_cell", "IAF ansvarar för tillsynen över a-kassorna.", None, 0),
        ]
        return structure

    def test_disposition_must_sit_on_a_body_section_heading(self):
        result = self.validate_in(
            self.structure_with_fact_box(),
            self.disposition(element_ids=["table_9_cell_1_1_p1", "paragraph_1"], quote="IAF:s tillsyn",
                             proposed_order=["Sammanfattning", "IAF:s tillsyn"]),
        )
        self.assertEqual(result.findings, [])
        self.assertEqual([r.reason for r in result.rejected], ["anchor_not_a_section_heading"])

    def test_heading_suggestions_may_still_sit_on_a_fact_box_heading(self):
        result = self.validate_in(
            self.structure_with_fact_box(),
            self.heading(element_ids=["table_9_cell_1_1_p1"], quote="IAF:s tillsyn"),
        )
        self.assertEqual(len(result.findings), 1, [r.reason for r in result.rejected])

    def test_proposed_order_may_only_name_body_section_headings(self):
        result = self.validate_in(
            self.structure_with_fact_box(),
            self.disposition(proposed_order=["IAF:s tillsyn", "Sammanfattning", "Resultat"]),
        )
        self.assertEqual(len(result.findings), 1)
        self.assertEqual(result.findings[0].proposed_order, [])

    def test_disposition_and_heading_are_limited_to_five_each(self):
        # Sju unika rubrikförslag: samma rubrik, olika kombinationer av relaterade ställen
        related = ["paragraph_2", "paragraph_4", "paragraph_6", "paragraph_8", "paragraph_9",
                   "table_1_cell_1_1_p1", "textbox_2_p1"]
        items = [self.heading(element_ids=["paragraph_3", other], description=f"Iakttagelse {i}.")
                 for i, other in enumerate(related)]
        result = self.validate(*items)
        self.assertEqual(len(result.findings), 5)
        self.assertEqual([r.reason for r in result.rejected], ["over_limit", "over_limit"])

    def test_same_places_in_two_categories_are_not_duplicates(self):
        result = self.validate(
            self.repetition(element_ids=["paragraph_9", "paragraph_8"],
                            quote="ledde 3,1 procent av tidrapporterna till återkrav"),
            self.inconsistency(),
        )
        self.assertEqual([f.category for f in result.findings], ["repetition", "inconsistency"])

    def test_unknown_ids_are_dropped_and_duplicates_removed(self):
        result = self.validate(
            self.repetition(element_ids=["paragraph_6", "paragraph_404", "paragraph_4"]),
            self.repetition(element_ids=["paragraph_4", "paragraph_6"]),
        )
        self.assertEqual(len(result.findings), 1)
        self.assertEqual(result.findings[0].element_ids, ["paragraph_6", "paragraph_4"])
        self.assertEqual([r.reason for r in result.rejected], ["duplicate"])


# ============================================================================
# Rendering som Word-kommentarer
# ============================================================================

def _body(docx_path):
    with ZipFile(docx_path) as z:
        return etree.fromstring(z.read("word/document.xml")).find(f"{W}body")


def _comments(docx_path):
    with ZipFile(docx_path) as z:
        root = etree.fromstring(z.read("word/comments.xml"))
    return {
        c.get(f"{W}id"): ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in c.findall(f"{W}p")]
        for c in root.findall(f"{W}comment")
    }


def _add_textbox(paragraph, text):
    """Minimal DrawingML-textruta i ett stycke (samma form som i Word)."""
    xml = (
        f'<w:r xmlns:w="{W_NS}" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<w:drawing><wp:anchor><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>' + text + '</w:t></w:r></w:p>'
        '</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>'
    )
    paragraph._p.append(etree.fromstring(xml))


class GlobalCommentRenderingTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"global-render-{id(self)}")
        self.source = self.root / "rapport.docx"

        document = Document()
        document.add_paragraph("Inledning", style="Heading 1")                          # paragraph_1
        document.add_paragraph("Granskningen omfattar alla a-kassor under 2024.")      # paragraph_2
        document.add_paragraph("Resultat", style="Heading 1")                           # paragraph_3
        document.add_paragraph("Granskningen omfattar alla a-kassor, som sagt.")        # paragraph_4
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Cellen upprepar samma sak."
        host = document.add_paragraph("Stycke med textruta.")                           # paragraph_5
        _add_textbox(host, "Faktaruta som upprepar avgränsningen.")
        document.save(self.source)
        self.structure = DocumentStructureExtractor(str(self.source), self.logger).extract()

    def tearDown(self):
        self.temp_dir.cleanup()

    def render(self, findings):
        output = self.root / "out.docx"
        with DocxPackage(str(self.source), self.logger) as pkg:
            results = GlobalFindingsRenderer(pkg, self.logger, self.structure).apply(findings)
            pkg.save(str(output))
        return output, results

    def finding(self, ids, **overrides):
        values = dict(category="repetition", element_ids=ids, quote="x",
                      description="Samma avgränsning står redan i inledningen.",
                      proposal="Stryk meningen här.")
        values.update(overrides)
        return GlobalFinding(**values)

    def test_comment_wraps_the_whole_paragraph_and_names_related_sections(self):
        output, results = self.render([self.finding(["paragraph_4", "paragraph_2"])])
        self.assertTrue(results[0].applied, results[0].message)

        paragraph = _body(output).findall(f"{W}p")[3]
        children = list(paragraph)
        start_index = next(i for i, c in enumerate(children) if c.tag == f"{W}commentRangeStart")
        self.assertTrue(all(c.tag == f"{W}pPr" for c in children[:start_index]))
        self.assertEqual(children[-2].tag, f"{W}commentRangeEnd")
        self.assertIsNotNone(children[-1].find(f"{W}commentReference"))

        lines = _comments(output)[str(results[0].comment_id)]
        self.assertEqual(lines[0], "Onödig upprepning. Samma avgränsning står redan i inledningen.")
        self.assertEqual(lines[1], "Förslag: Stryk meningen här.")
        self.assertEqual(
            lines[2],
            "Hänger ihop med: avsnittet ”Inledning” (”Granskningen omfattar alla a-kassor under 2024.”).",
        )
        with ZipFile(output) as z:
            comment = etree.fromstring(z.read("word/comments.xml")).find(f"{W}comment")
        self.assertEqual(comment.get(f"{W}author"), "JBG Klarspråkningstjänst (global granskning)")
        Document(str(output))

    def test_inconsistency_comment_quotes_the_conflicting_statement(self):
        output, results = self.render([self.finding(
            ["paragraph_4", "paragraph_2"], category="inconsistency",
            description="Avgränsningen beskrivs olika i två avsnitt.",
            proposal="Använd samma avgränsning på båda ställena.",
            related_quote="alla a-kassor under 2024",
        )])
        self.assertTrue(results[0].applied, results[0].message)
        lines = _comments(output)[str(results[0].comment_id)]
        self.assertEqual(lines, [
            "Inkonsekvent påstående. Avgränsningen beskrivs olika i två avsnitt.",
            "Förslag: Använd samma avgränsning på båda ställena.",
            "Jämför med: avsnittet ”Inledning” (”alla a-kassor under 2024”).",
        ])

    def test_error_comment_without_related_places_has_no_related_line(self):
        output, results = self.render([self.finding(
            ["paragraph_4"], category="error",
            description="Stycket innehåller kvarlämnad arbetstext.",
            proposal="Ta bort ”som sagt”.",
        )])
        self.assertTrue(results[0].applied, results[0].message)
        self.assertEqual(_comments(output)[str(results[0].comment_id)], [
            "Troligt fel. Stycket innehåller kvarlämnad arbetstext.",
            "Förslag: Ta bort ”som sagt”.",
        ])

    def test_disposition_comment_shows_the_proposed_order(self):
        output, results = self.render([self.finding(
            ["paragraph_3", "paragraph_1"], category="disposition", quote="Resultat",
            description="Resultaten presenteras innan läsaren vet hur de har tagits fram.",
            proposal="Överväg att låta metoden komma före resultaten.",
            proposed_order=["Inledning", "Resultat"],
        )])
        self.assertTrue(results[0].applied, results[0].message)
        self.assertEqual(_comments(output)[str(results[0].comment_id)], [
            "Förslag om disposition. Resultaten presenteras innan läsaren vet hur de har tagits fram.",
            "Förslag: Överväg att låta metoden komma före resultaten.",
            "Se även: rubriken ”Inledning”.",
            "Möjlig ordning: ”Inledning”, ”Resultat”.",
        ])
        # Kommentaren sitter på rubriken
        heading = _body(output).findall(f"{W}p")[2]
        self.assertEqual("".join(t.text or "" for t in heading.iter(f"{W}t")), "Resultat")
        self.assertIsNotNone(heading.find(f"{W}commentRangeStart"))

    def test_table_cell_and_textbox_anchors(self):
        textbox_id = next(e["element_id"] for e in self.structure["elements"] if e["type"] == "textbox")
        output, results = self.render([
            self.finding(["table_1_cell_1_1_p1", "paragraph_2"]),
            self.finding([textbox_id, "paragraph_2"]),
        ])
        self.assertTrue(all(r.applied for r in results), [r.message for r in results])
        body = _body(output)
        cell_paragraph = body.find(f"{W}tbl").find(f".//{W}p")
        self.assertIsNotNone(cell_paragraph.find(f"{W}commentRangeStart"))
        # Textrutan förankras i stycket som bär den, inte inne i textrutan
        host = body.findall(f"{W}p")[4]
        self.assertIsNotNone(host.find(f"{W}commentRangeStart"))
        self.assertIsNone(host.find(f".//{W}txbxContent//{W}commentRangeStart"))
        Document(str(output))

    def test_unsupported_anchor_is_skipped_without_failing_the_rest(self):
        self.structure["elements"].append(
            {"type": "footnote", "element_id": "footnote_1", "text": "Fotnot", "doc_order": None}
        )
        output, results = self.render([
            self.finding(["footnote_1", "paragraph_2"]),
            self.finding(["paragraph_4", "paragraph_2"]),
        ])
        self.assertEqual([r.applied for r in results], [False, True])
        self.assertEqual(len(_comments(output)), 1)


# ============================================================================
# Hela kedjan
# ============================================================================

class _FakeCompletions:
    """Svarar olika på lokala och globala anrop (känns igen på systemprompten)."""

    def __init__(self, global_content):
        self.global_content = global_content
        self.global_calls = 0

    def create(self, **kwargs):
        is_global = "dokumentnivå" in kwargs["messages"][0]["content"]
        if is_global:
            self.global_calls += 1
            content = self.global_content
            usage = SimpleNamespace(prompt_tokens=3000, completion_tokens=400, total_tokens=3400)
        else:
            content = "[]"
            usage = SimpleNamespace(prompt_tokens=1000, completion_tokens=50, total_tokens=1050)
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))], usage=usage,
        )


class GlobalReviewEndToEndTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"global-e2e-{id(self)}")
        self.source = self.root / "rapport.docx"
        document = Document()
        document.add_paragraph("Inledning", style="Heading 1")
        document.add_paragraph("Granskningen omfattar alla a-kassor under 2024.")
        document.add_paragraph("Resultat", style="Heading 1")
        document.add_paragraph("Granskningen omfattar alla a-kassor under 2024, som nämnts.")
        document.save(self.source)

    def tearDown(self):
        self.temp_dir.cleanup()

    def run_improver(self, global_content, global_review=True, docx_mode="tracked"):
        completions = _FakeCompletions(global_content)
        client = SimpleNamespace(chat=SimpleNamespace(completions=completions))
        output = self.root / f"out_{docx_mode}.docx"
        with mock.patch.object(suggestor_module.openai, "OpenAI", return_value=client), \
             mock.patch.object(analyzer_module.openai, "OpenAI", return_value=client):
            improver = JBGLanguageImprover(
                input_path=str(self.source), api_key="unused", model="test-model",
                prompt_policy="lokal policy", temperature=1, include_motivations=True,
                logger=self.logger, docx_mode=docx_mode, global_review=global_review,
            )
            improver.run(output_path=str(output))
        return improver, output, completions

    FINDING = json.dumps({"findings": [{
        "category": "repetition",
        "element_ids": ["paragraph_4", "paragraph_2"],
        "quote": "Granskningen omfattar alla a-kassor under 2024, som nämnts.",
        "description": "Samma avgränsning står redan i inledningen.",
        "proposal": "Stryk meningen.",
    }]}, ensure_ascii=False)

    def test_findings_become_comments_and_are_reported(self):
        for mode in ("tracked", "simple"):
            with self.subTest(mode=mode):
                improver, output, completions = self.run_improver(self.FINDING, docx_mode=mode)
                self.assertEqual(completions.global_calls, 1)

                comments = [" ".join(lines) for lines in _comments(output).values()]
                self.assertTrue(any(c.startswith("Onödig upprepning. ") for c in comments))

                summary = improver.run_summary
                self.assertEqual(summary.global_findings["accepted"], 1)
                self.assertEqual(summary.global_findings["comments_applied"], 1)
                self.assertEqual(summary.global_findings["by_category"], {"repetition": 1})
                self.assertEqual(summary.usage["by_phase"]["global"]["prompt_tokens"], 3000)

                tables = about_tables(_body(output))
                self.assertIn([["Kategori", "Antal"], ["Upprepningar", "1"], ["Totalt", "1"]], tables)
                usage = next(t for t in tables if t[0][0] == "")
                self.assertEqual(usage[0], ["", "Lokal granskning", "Global granskning", "Totalt"])
                self.assertEqual(usage[1], ["Anrop till språkmodellen", "1", "1", "2"])

                saved = json.loads(Path(improver.global_findings_json).read_text(encoding="utf-8"))
                self.assertEqual(saved["summary"]["accepted"], 1)

    def test_both_categories_come_from_one_call_and_are_counted(self):
        content = json.dumps({"findings": json.loads(self.FINDING)["findings"] + [{
            "category": "inconsistency",
            "element_ids": ["paragraph_4", "paragraph_2"],
            "quote": "under 2024, som nämnts",
            "related_quote": "Granskningen omfattar alla a-kassor under 2024.",
            "description": "Perioden beskrivs olika.",
            "proposal": "Kontrollera perioden.",
        }, {
            "category": "error",
            "element_ids": ["paragraph_2"],
            "quote": "Granskningen omfattar alla a-kassor under 2024.",
            "description": "Meningen saknar avslutning.",
            "proposal": "Komplettera meningen.",
        }]}, ensure_ascii=False)
        improver, output, completions = self.run_improver(content, docx_mode="simple")
        self.assertEqual(completions.global_calls, 1)
        self.assertEqual(improver.run_summary.global_findings["by_category"],
                         {"repetition": 1, "inconsistency": 1, "error": 1})
        comments = [" ".join(lines) for lines in _comments(output).values()]
        self.assertTrue(any(c.startswith("Inkonsekvent påstående. Perioden beskrivs olika.") for c in comments))
        self.assertIn(
            [["Kategori", "Antal"], ["Upprepningar", "1"], ["Inkonsekvenser", "1"],
             ["Troliga fel", "1"], ["Totalt", "3"]],
            about_tables(_body(output)),
        )

    def test_no_global_call_when_unchecked(self):
        improver, _, completions = self.run_improver(self.FINDING, global_review=False)
        self.assertEqual(completions.global_calls, 0)
        self.assertIsNone(improver.run_summary.global_findings)
        self.assertNotIn("global", improver.run_summary.usage["by_phase"])

    def test_a_broken_global_answer_never_breaks_the_run(self):
        improver, output, _ = self.run_improver("Det här är inte JSON.")
        self.assertTrue(improver.run_summary.succeeded)
        self.assertEqual(improver.run_summary.global_findings["accepted"], 0)
        self.assertTrue(improver.run_summary.global_findings["errors"])
        texts = ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in _body(output).findall(f"{W}p")]
        self.assertIn("Den globala granskningen kunde inte genomföras.", texts)

    def test_no_findings_is_reported_as_zero(self):
        improver, output, _ = self.run_improver('{"findings": []}')
        self.assertIn([["Kategori", "Antal"], ["Totalt", "0"]], about_tables(_body(output)))


if __name__ == "__main__":
    unittest.main()
