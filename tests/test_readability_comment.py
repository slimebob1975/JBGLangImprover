import json
import logging
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from zipfile import ZipFile

from docx import Document
from lxml import etree

from app.src import JBGLangImprovSuggestorAI as suggestor_module
from app.src.JBGLanguageImprover import JBGLanguageImprover
from app.src.JBGReadabilityCommentRenderer import build_readability_comment, find_title_element


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"


def _quiet_logger(name):
    logger = logging.getLogger(name)
    logger.handlers.clear()
    logger.addHandler(logging.NullHandler())
    logger.propagate = False
    return logger


def _el(eid, etype, text, level=None, order=None):
    return {"type": etype, "element_id": eid, "text": text, "heading_level": level, "doc_order": order}


class TitleTests(unittest.TestCase):
    def test_title_style_is_preferred(self):
        structure = {"elements": [
            _el("paragraph_1", "paragraph", "Omslagstext", None, 1),
            _el("paragraph_2", "paragraph", "Rapportens titel", 0, 2),
        ]}
        self.assertEqual(find_title_element(structure)["element_id"], "paragraph_2")

    def test_first_body_paragraph_with_text_when_no_title_style(self):
        # Testkörningen b391e3ac: titeln står i Normal före faktarutan och förordet
        structure = {"elements": [
            _el("paragraph_1", "paragraph", "Mörkertal för felaktiga utbetalningar", None, 1),
            _el("table_1_cell_1_1_p1", "table_cell", "IAF:s tillsyn", 1, 2),
            _el("paragraph_3", "paragraph", "Förord", 1, 7),
        ]}
        self.assertEqual(find_title_element(structure)["element_id"], "paragraph_1")

    def test_fact_boxes_and_empty_paragraphs_are_skipped(self):
        structure = {"elements": [
            _el("table_1_cell_1_1_p1", "table_cell", "Faktaruta", 1, 1),
            _el("paragraph_1", "paragraph", "   ", None, 2),
            _el("paragraph_2", "paragraph", "Inledning", 1, 3),
        ]}
        self.assertEqual(find_title_element(structure)["element_id"], "paragraph_2")

    def test_comment_text(self):
        text = build_readability_comment({
            "before": {"lix": 53.4, "band": "Svår"}, "after": {"lix": 52.4, "band": "Svår"},
        })
        lines = text.split("\n")
        self.assertEqual(lines[0], "Läsbarhet (LIX): 53,4 (svår) före granskningen och 52,4 (svår) "
                                   "om alla förslag godtas, en förändring med \u22121,0.")
        self.assertIn("Riktvärden", lines[2])
        unchanged = build_readability_comment({"before": {"lix": 40.0}, "after": {"lix": 40.04}})
        self.assertIn("om alla förslag godtas, oförändrat.", unchanged)


class _FakeCompletions:
    def create(self, **kwargs):
        content = json.dumps([{
            "type": "paragraph", "element_id": "paragraph_3",
            "old": "genomfört en granskning av", "new": "granskat",
            "motivation": "Verb i stället för substantiv.",
        }], ensure_ascii=False)
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))],
            usage=SimpleNamespace(prompt_tokens=100, completion_tokens=10, total_tokens=110),
        )


class ReadabilityCommentEndToEndTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.logger = _quiet_logger(f"lix-comment-{id(self)}")
        self.source = self.root / "rapport.docx"
        document = Document()
        document.add_paragraph("Mörkertal för felaktiga utbetalningar")      # titel i Normal
        document.add_paragraph("Inledning", style="Heading 1")
        document.add_paragraph("Vi har genomfört en granskning av ärendet.")
        document.save(self.source)

    def tearDown(self):
        self.temp_dir.cleanup()

    def run_improver(self, include_about_section, compute_readability, docx_mode="simple"):
        client = SimpleNamespace(chat=SimpleNamespace(completions=_FakeCompletions()))
        output = self.root / f"out_{include_about_section}_{compute_readability}_{docx_mode}.docx"
        with mock.patch.object(suggestor_module.openai, "OpenAI", return_value=client):
            improver = JBGLanguageImprover(
                input_path=str(self.source), api_key="unused", model="m", prompt_policy="p",
                temperature=1, include_motivations=False, logger=self.logger, docx_mode=docx_mode,
                include_about_section=include_about_section, compute_readability=compute_readability,
            )
            improver.run(output_path=str(output))
        with ZipFile(output) as z:
            body = etree.fromstring(z.read("word/document.xml")).find(f"{W}body")
            comments = (
                [" ".join("".join(t.text or "" for t in p.iter(f"{W}t")) for p in c.findall(f"{W}p"))
                 for c in etree.fromstring(z.read("word/comments.xml")).findall(f"{W}comment")]
                if "word/comments.xml" in z.namelist() else []
            )
        texts = ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in body.findall(f"{W}p")]
        return improver, body, texts, comments

    def test_lix_becomes_a_comment_on_the_title_without_the_section(self):
        for mode in ("simple", "tracked"):
            with self.subTest(mode=mode):
                improver, body, texts, comments = self.run_improver(False, True, mode)
                self.assertNotIn("Om klarspråkningen", texts)
                lix = [c for c in comments if c.startswith("Läsbarhet (LIX):")]
                self.assertEqual(len(lix), 1)
                self.assertEqual(improver.run_summary.lix_comment["element_id"], "paragraph_1")
                title = body.findall(f"{W}p")[0]
                self.assertIsNotNone(title.find(f"{W}commentRangeStart"))

    def test_no_lix_comment_when_the_section_shows_lix(self):
        improver, _, texts, comments = self.run_improver(True, True)
        self.assertIn("Om klarspråkningen", texts)
        self.assertFalse(any(c.startswith("Läsbarhet (LIX):") for c in comments))
        self.assertIsNone(improver.run_summary.lix_comment)

    def test_nothing_when_lix_is_off(self):
        improver, _, texts, comments = self.run_improver(False, False)
        self.assertNotIn("Om klarspråkningen", texts)
        self.assertFalse(any(c.startswith("Läsbarhet (LIX):") for c in comments))
        self.assertIsNone(improver.run_summary.lix_comment)
        self.assertIsNone(improver.run_summary.readability)


if __name__ == "__main__":
    unittest.main()
