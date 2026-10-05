import json
import logging
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from app.src.JBGDocumentStructureExtractor import DocumentStructureExtractor
from app.src.JBGGlobalAnalyzerAI import GlobalReviewResult, JBGGlobalAnalyzerAI
from app.src.JBGLangImprovSuggestorAI import JBGLangImprovSuggestorAI


def _quiet_logger(name):
    logger = logging.getLogger(name)
    logger.handlers.clear()
    logger.addHandler(logging.NullHandler())
    logger.propagate = False
    return logger


def _el(eid, etype, text, level=None, order=None):
    return {"type": etype, "element_id": eid, "text": text, "heading_level": level,
            "doc_order": order, "style_id": "Normal", "style_name": "Normal"}


class HeadingStyleNameTests(unittest.TestCase):
    """Testkörningen c7c86075: "IAF Rubrik 3 numrerad" saknar dispositionsnivå."""

    def test_heading_number_anywhere_in_the_style_name(self):
        with tempfile.TemporaryDirectory() as tmp:
            document = Document()
            for name in ("IAF Rubrik 3 numrerad", "IAF Rubrik till textruta 2", "IAF Bilagerubrik 1"):
                document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)   # utan outlineLvl
            document.add_paragraph("Programanvisningen ska återkallas vid misskötsel.", style="IAF Rubrik 3 numrerad")
            document.add_paragraph("Faktaruta", style="IAF Rubrik till textruta 2")
            document.add_paragraph("Bilaga", style="IAF Bilagerubrik 1")
            path = Path(tmp) / "d.docx"
            document.save(path)
            levels = {e["element_id"]: e["heading_level"]
                      for e in DocumentStructureExtractor(str(path), _quiet_logger("h")).extract()["elements"]}
        self.assertEqual(levels, {"paragraph_1": 3, "paragraph_2": None, "paragraph_3": None})


class FinalHeadingPunctuationTests(unittest.TestCase):
    """Policyn: rubriker avslutas aldrig med punkt. Borttagningen ska inte stoppas."""

    def accepted(self, text, old, new, level):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "s.json"
            path.write_text(json.dumps({"type": "docx", "elements": [
                _el("paragraph_1", "paragraph", text, level, 1)]}, ensure_ascii=False), encoding="utf-8")
            suggestor = JBGLangImprovSuggestorAI(api_key="k", model="m", prompt_policy="p",
                                                 temperature=1, logger=_quiet_logger("p"))
            suggestor.load_structure(str(path))
            return bool(suggestor._postprocess_suggestions([
                {"type": "paragraph", "element_id": "paragraph_1", "old": old, "new": new}]))

    def test_final_period_or_colon_removed_from_a_heading(self):
        self.assertTrue(self.accepted("Återkallelsegrunderna är otydligt beskrivna.", ".", "", 2))
        self.assertTrue(self.accepted("Bilaga 2:", ":", "", 1))

    def test_not_in_body_text(self):
        self.assertFalse(self.accepted("Detta är en vanlig mening.", ".", "", None))

    def test_not_when_the_heading_contains_the_mark_twice(self):
        # Ett ensamt "." kan inte visa vilket av tecknen som avses
        self.assertFalse(self.accepted("Kap. 3 Resultat.", ".", "", 1))


class CoverPlaceholderAndHeadingRepetitionTests(unittest.TestCase):
    STRUCTURE = {"type": "docx", "elements": [
        _el("paragraph_1", "paragraph", "Rapportens titel", None, 1),
        _el("paragraph_2", "paragraph", "Utgiven: Månad ÅÅÅÅ https://www.iaf.se", None, 2),
        _el("paragraph_3", "paragraph", "Diarienummer: IAF 2026/78", None, 3),
        _el("paragraph_4", "paragraph", "Sammanfattning", 1, 4),
        _el("table_1_cell_1_1_p1", "table_cell", "IAF riktar en anmärkning till Arbetsförmedlingen för att", 1, 5),
        _el("table_1_cell_1_1_p2", "table_cell", "Beslut saknar rättslig grund.", None, 6),
        _el("paragraph_5", "paragraph", "Slutsatser", 1, 7),
        _el("table_2_cell_1_1_p1", "table_cell", "IAF riktar en anmärkning till Arbetsförmedlingen för att", 1, 8),
        _el("table_2_cell_1_1_p2", "table_cell", "Kommuniceringen brister.", None, 9),
        _el("paragraph_6", "paragraph", "IAF riktar en anmärkning till Arbetsförmedlingen för att", None, 10),
    ]}

    def validate(self, *items):
        result = GlobalReviewResult()
        JBGGlobalAnalyzerAI("k", "m", 1, _quiet_logger("g"), policy="p").validate(
            self.STRUCTURE, [(1, item) for item in items], result)
        return result

    def test_unfilled_placeholder_on_the_cover_is_reported(self):
        # Testkörningen c7c86075: "Utgiven: Månad ÅÅÅÅ" avvisades som omslag
        result = self.validate({"category": "error", "element_ids": ["paragraph_2"],
                                "quote": "Utgiven: Månad ÅÅÅÅ", "description": "Platshållare kvar."})
        self.assertEqual(len(result.findings), 1, [r.reason for r in result.rejected])

    def test_other_cover_findings_are_still_rejected(self):
        result = self.validate(
            {"category": "error", "element_ids": ["paragraph_3"], "quote": "Diarienummer: IAF 2026/78",
             "description": "Ser ofullständigt ut."},
            {"category": "heading", "element_ids": ["paragraph_1"], "quote": "Rapportens titel",
             "description": "Rubriken kan bli tydligare."},
        )
        self.assertEqual([r.reason for r in result.rejected], ["cover_material", "cover_material"])

    def test_repeated_headings_are_not_a_repetition(self):
        # Testkörningen c7c86075: samma inledning i flera faktarutor
        result = self.validate({"category": "repetition",
                                "element_ids": ["table_2_cell_1_1_p1", "table_1_cell_1_1_p1"],
                                "quote": "IAF riktar en anmärkning till Arbetsförmedlingen för att",
                                "description": "Samma ruta återkommer."})
        self.assertEqual([r.reason for r in result.rejected], ["repetition_of_headings"])

    def test_a_heading_repeated_in_body_text_can_still_be_a_repetition(self):
        result = self.validate({"category": "repetition",
                                "element_ids": ["paragraph_6", "table_1_cell_1_1_p1"],
                                "quote": "IAF riktar en anmärkning till Arbetsförmedlingen för att",
                                "description": "Rutans rubrik upprepas i löptexten."})
        self.assertEqual(len(result.findings), 1, [r.reason for r in result.rejected])


if __name__ == "__main__":
    unittest.main()
